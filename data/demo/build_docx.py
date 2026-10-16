"""Build two small, synthetic import fixtures; contains no real personal information.

Run with the document-tool bundled Python or an environment with python-docx.
This generates source materials for import, never a finished business report.
"""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent


def build(path: Path, title: str, paragraphs: list[str], rows: list[tuple[str, str]]) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ("Normal", "Title", "Heading 1"):
        style = document.styles[name]
        style.font.name = "Arial Unicode MS"
        style.font.color.rgb = RGBColor(0, 0, 0)
        properties = style.element.get_or_add_pPr()
        for border in properties.findall(qn("w:pBdr")):
            properties.remove(border)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{script}"), "Arial Unicode MS")
        for attribute in tuple(fonts.attrib):
            if attribute.lower().endswith("theme"):
                del fonts.attrib[attribute]
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.space_after = Pt(8)
    document.styles["Title"].font.size = Pt(20)
    document.add_paragraph(title, "Title")
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "事项"
    table.rows[0].cells[1].text = "模拟记录"
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for attribute, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{attribute}"), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for cell in table.rows[0].cells:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E6EEF7")
        cell._tc.get_or_add_tcPr().append(shading)
    for label, content in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = content
    document.add_paragraph("所有内容均为虚构，仅用于材料导入、段落及表格解析演示。")
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def main() -> None:
    build(
        ROOT / "technical-q2" / "季度工作记录.docx",
        "技术部季度工作记录 模拟材料",
        [
            "用途：2026年第二季度部门汇报的文字材料。项目、人员编号与事件均为虚构。",
            "统计范围为2026年4月1日至6月30日。具体项目状态和成果以同目录CSV为准，本文仅补充背景与待确认叙述。",
            "考勤仅包含两个模拟排班日，不应推广成全季度出勤表现。",
        ],
        [
            ("工作背景", "模拟开展内部知识库与审批表单改造，梳理接口文档及监控需求。"),
            ("待确认问题", "监控需求口径仍待确认；该叙述不表示已经发生的真实故障。"),
            ("下期建议", "建议补充验收记录，明确接口文档责任范围并完善监控需求清单。"),
        ],
    )
    build(
        ROOT / "administration-q1" / "服务记录.docx",
        "行政服务记录 模拟材料",
        [
            "用途：2026年第一季度行政部满意度汇报的文字材料，不对应任何真实部门。",
            "同目录CSV包含4份虚构反馈，使用统一1—5分量表。样本数量有限，不能代表全体员工。",
            "没有往期对照数据，不能由这些记录推断同比或环比变化。",
        ],
        [
            ("服务范围", "模拟登记咨询、会议支持和物资申请等日常服务。"),
            ("改进建议", "建议统一申请指引，补充处理进度反馈；这些是建议，尚非已完成成果。"),
            ("证据限制", "不包含实际服务总量、响应时长或投诉率，不应虚构这些指标。"),
        ],
    )


if __name__ == "__main__":
    main()
