"""Build the editable course report from reviewed Markdown and actual figures.

Optional dependencies: uv sync --locked --group report.
This builder never loads model configuration or calls a model service.
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "docs" / "report"


def font(style, size, *, bold=False):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "PingFang SC")


def inline(paragraph, text):
    # Keep link labels and their real targets readable in a standalone Word copy.
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        run = paragraph.add_run(token.strip("*`") if token.startswith(("**", "`")) else token)
        if token.startswith("**"):
            run.bold = True
        elif token.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(10)


def table(doc, rows):
    widths = {
        2: [1.65, 5.4],
        3: [1.15, 2.0, 3.9],
        4: [1.15, 1.35, 1.35, 3.2],
        5: [1.35, 1.3, 1.15, 1.1, 2.15],
        6: [1.25, 1.15, 0.85, 0.85, 1.0, 1.95],
    }.get(len(rows[0]), [7.05 / len(rows[0])] * len(rows[0]))
    item = doc.add_table(rows=0, cols=len(rows[0]))
    item.autofit = False
    for col, width in zip(item.columns, widths, strict=True):
        col.width = Inches(width)
    for index, values in enumerate(rows):
        row = item.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            tr_pr.append(repeat)
        for col, (cell, width) in enumerate(zip(row.cells, widths, strict=True)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.keep_with_next = False
            value = values[col] if col < len(values) else ""
            inline(p, value)
            if len(value) <= 16:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(10.5)
                if index == 0:
                    run.bold = True
            properties = cell._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "DCE6EF" if index == 0 else "FFFFFF")
            properties.append(shading)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{edge}")
                for attr, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn(f"w:{attr}"), val)
                borders.append(border)
            properties.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                child = OxmlElement(f"w:{edge}")
                child.set(qn("w:w"), "100")
                child.set(qn("w:type"), "dxa")
                margins.append(child)
            properties.append(margins)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(doc, path, caption):
    if not path.is_file():
        raise FileNotFoundError(f"Required reviewed figure is missing: {path.name}")
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    picture = run.add_picture(str(path), width=Inches(6.95))
    picture._inline.docPr.set("descr", caption)
    p = doc.add_paragraph(caption, style="Caption")
    p.paragraph_format.keep_with_next = False


def markdown(doc, source, *, skip_title=False):
    lines = source.read_text(encoding="utf-8").splitlines()
    if skip_title and source.name == "experiment-report.md":
        # The Word cover has its own editable identity fields. Start the body
        # at the first section instead of duplicating Markdown cover metadata.
        lines = lines[next(i for i, row in enumerate(lines) if row.startswith("## ")) :]
    index = 0
    while index < len(lines):
        line = lines[index]
        index += 1
        if not line.strip():
            continue
        if line.startswith("```"):
            language = line.removeprefix("```").strip()
            block = []
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            if language == "mermaid":
                # The Markdown retains the editable graph source; Word uses the
                # reviewed repository architecture figure rather than raw syntax.
                diagram = (
                    "sequence.png"
                    if any("sequenceDiagram" in row for row in block)
                    else "architecture.png"
                )
                caption = (
                    "实际协作与审核返工时序"
                    if diagram == "sequence.png"
                    else "实际系统架构与有界协作流程"
                )
                figure(doc, REPORT_DIR / "figures" / diagram, caption)
            else:
                for code in block:
                    p = doc.add_paragraph(code, style="Code")
                    p.paragraph_format.keep_with_next = False
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
        if image_match:
            figure(doc, (source.parent / image_match[2]).resolve(), image_match[1])
            continue
        if line.strip().startswith("|"):
            rows = []
            while True:
                values = [part.strip() for part in line.strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", value) for value in values):
                    rows.append(values)
                if index >= len(lines) or not lines[index].strip().startswith("|"):
                    break
                line = lines[index]
                index += 1
            table(doc, rows)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            if skip_title and len(heading[1]) == 1:
                continue
            title = heading[2].replace("**", "").replace("`", "")
            title = re.sub(r"^([一二三四五六七八九十]+)[、.]", r"\1 ", title)
            title = re.sub(r"^(\d+)[.、]\s*", r"\1 ", title)
            title = re.sub(r"[^\w\s\u4e00-\u9fff]", " ", title)
            level = min(len(heading[1]) - 1, 3)
            doc.add_heading(title, level=max(level, 1))
            continue
        style = "List Bullet" if line.startswith("- ") else "Normal"
        if style == "List Bullet":
            line = line[2:]
        if line.startswith("> "):
            line = line[2:]
        inline(doc.add_paragraph(style=style), line)


def build(output):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.72)
    font(doc.styles["Normal"], 11.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.22
    for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12)):
        font(doc.styles[name], size, bold=True)
        doc.styles[name].paragraph_format.space_before = Pt(14)
        doc.styles[name].paragraph_format.space_after = Pt(7)
    font(doc.styles["Caption"], 10)
    doc.styles["Caption"].paragraph_format.space_after = Pt(9)
    code_style = doc.styles.add_style("Code", 1)
    font(code_style, 9)
    code_style.font.name = "Consolas"
    code_style.paragraph_format.line_spacing = 1.1
    code_style.paragraph_format.space_after = Pt(4)
    doc.add_paragraph("多智能体办公协作系统实验报告", style="Title")
    doc.add_paragraph("软件体系结构研究生课程  LangGraph 与 Streamlit  v1.0.0")
    doc.add_paragraph("姓名 __________________  学号 __________________")
    doc.add_paragraph("班级 __________________  日期 2026年10月10日")
    doc.add_paragraph("代码仓库 https://github.com/Thedaiyyds/multi-agent-office-lab")
    markdown(doc, REPORT_DIR / "experiment-report.md", skip_title=True)
    doc.add_page_break()
    doc.add_heading("附录 体系结构分析", level=1)
    markdown(doc, REPORT_DIR / "architecture-analysis.md", skip_title=True)
    doc.core_properties.title = "多智能体办公协作系统实验报告"
    doc.core_properties.subject = "软件体系结构课程实验"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Editable report built: {output}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=REPORT_DIR / "experiment-report.docx")
    args = parser.parse_args()
    build(args.output.resolve())


if __name__ == "__main__":
    main()
