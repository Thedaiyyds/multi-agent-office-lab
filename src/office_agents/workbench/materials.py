"""Bounded local material imports, validated before immutable confirmation."""

import csv
import hashlib
import io
import json
import math
import re
import stat
import zipfile
from datetime import UTC, date, datetime
from pathlib import Path, PurePosixPath
from typing import Annotated, Literal
from uuid import uuid4

from docx import Document
from fastapi import APIRouter, File, Form, Header, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse
from lxml import etree
from pydantic import BaseModel, ConfigDict, Field

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
MAX_DOCX_BYTES = 50 * 1024 * 1024
Kind = Literal["projects", "attendance", "achievements", "satisfaction", "notes"]
Workspace = Literal["normal", "demo"]
FIELDS = {
    "projects": [
        "project_id",
        "department",
        "effective_date",
        "snapshot_date",
        "status",
        "progress",
    ],
    "attendance": ["person_id", "department", "date", "planned_hours", "actual_hours"],
    "achievements": ["achievement_id", "department", "date", "category", "name"],
    "satisfaction": ["feedback_id", "department", "date", "score", "scale_min", "scale_max"],
    "notes": [],
}
OPTIONAL_FIELDS = {
    "projects": ["completed_date"],
    "attendance": ["leave_hours", "absent_hours", "late_minutes"],
    "achievements": ["project_id"],
    "satisfaction": ["valid"],
    "notes": [],
}
ALIASES = {
    "project_id": ["项目id", "项目编号"],
    "department": ["部门", "归属部门", "服务部门"],
    "effective_date": ["生效日期", "立项日期"],
    "snapshot_date": ["快照日期", "统计日期"],
    "status": ["状态", "项目状态"],
    "progress": ["进度", "项目进度"],
    "person_id": ["人员id", "人员编号", "员工编号"],
    "date": ["日期", "发生日期"],
    "planned_hours": ["计划工时", "计划出勤工时"],
    "actual_hours": ["实际工时", "实际出勤工时"],
    "achievement_id": ["成果id", "成果编号"],
    "category": ["类别", "成果类别"],
    "name": ["名称", "成果名称"],
    "feedback_id": ["反馈id", "反馈编号"],
    "score": ["评分", "满意度评分"],
    "scale_min": ["量表下限", "最低分"],
    "scale_max": ["量表上限", "最高分"],
    "completed_date": ["实际完成日期", "完成日期"],
    "leave_hours": ["请假工时", "请假小时"],
    "absent_hours": ["缺勤工时", "缺勤小时"],
    "late_minutes": ["迟到分钟", "迟到时长"],
    "valid": ["有效标记", "是否有效"],
}


class ConfirmInput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    mapping: dict[str, str] = Field(default_factory=dict, max_length=100)
    department: str | None = Field(default=None, max_length=80)
    start_date: str | None = Field(default=None, max_length=10)
    end_date: str | None = Field(default=None, max_length=10)


def issue(message, location="文件", severity="error"):
    return {"severity": severity, "location": location, "message": message}


def reject(message, location="文件", status=422):
    raise HTTPException(status, {"message": message, "issues": [issue(message, location)]})


def canonical_json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def valid_date(value):
    return (
        isinstance(value, str)
        and bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", value))
        and date.fromisoformat(value).isoformat() == value
    )


def decode_text(data):
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = data.decode(encoding)
            if "\x00" in text or any(ord(c) < 32 and c not in "\t\r\n" for c in text):
                reject("文件包含二进制或控制字符，请提供文本文件")
            return text
        except UnicodeDecodeError:
            continue
    reject("无法无损解码，请另存为 UTF-8 或 GB18030")


def table_rows(entries):
    """Turn location-tagged cells into records; preserve physical locations."""
    if not entries:
        reject("材料为空")
    columns = [str(cell).strip() for cell in entries[0][1]]
    if not columns or any(not c for c in columns) or len(columns) != len(set(columns)):
        reject("表头不能为空或重复", entries[0][0])
    if len(columns) > 100 or any(len(c) > 200 for c in columns):
        reject("表头最多 100 列，每列名称不超过 200 字")
    rows, issues = [], []
    for location, cells in entries[1:]:
        if not any(str(c).strip() for c in cells):
            continue
        if len(cells) != len(columns):
            issues.append(issue("列数与表头不一致", location))
        rows.append(
            {"location": location, "values": dict(zip(columns, [str(c).strip() for c in cells]))}
        )
        if len(rows) > MAX_ROWS:
            reject("明细超过 20000 行，请拆分文件")
    if not rows:
        issues.append(issue("表格没有数据行"))
    return columns, rows, issues


def parse_csv(data):
    reader = csv.reader(io.StringIO(decode_text(data), newline=""), strict=True)
    entries = []
    try:
        while True:
            start = reader.line_num + 1
            cells = next(reader)
            end = reader.line_num
            if not any(c.strip() for c in cells):
                continue
            if len(entries) > MAX_ROWS:
                reject("明细超过 20000 行，请拆分文件")
            location = f"行 {start}" if start == end else f"行 {start}–{end}"
            entries.append((location, cells))
    except StopIteration:
        pass
    except (csv.Error, UnicodeError):
        reject("CSV 格式错误，请检查引号、分隔符和字段长度", f"行 {reader.line_num}")
    return table_rows(entries)


def parse_docx(data, kind):
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > 2000 or len(names) != len(set(names)):
                reject("DOCX 压缩包条目过多或重复")
            if sum(entry.file_size for entry in entries) > MAX_DOCX_BYTES:
                reject("DOCX 解压后超过 50MiB，请拆分文件")
            if not {"[Content_Types].xml", "word/document.xml", "_rels/.rels"}.issubset(names):
                reject("文件不是真实的 DOCX 文档")
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if (
                    path.is_absolute()
                    or ".." in path.parts
                    or "\\" in entry.filename
                    or ":" in entry.filename
                    or entry.flag_bits & 1
                    or stat.S_ISLNK(entry.external_attr >> 16)
                ):
                    reject("DOCX 包含不安全的压缩条目")
                lowered = entry.filename.lower()
                if "vbaproject" in lowered or "activex/" in lowered or "embeddings/" in lowered:
                    reject("不支持 DOCX 宏、嵌入文件或活动对象")
                if lowered.endswith((".xml", ".rels")):
                    content = archive.read(entry)
                    if b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper():
                        reject("DOCX 包含不支持的 XML 声明")
                    xml = etree.fromstring(content, parser)
                    if xml.xpath("//*[local-name()='Relationship' and @TargetMode='External']"):
                        reject("DOCX 包含外部链接，请移除链接后导入")
                    if lowered == "[content_types].xml":
                        types = xml.xpath("//@ContentType")
                        if (
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
                            not in types
                            or any("macroenabled" in value.lower() for value in types)
                        ):
                            reject("文件类型不是无宏 DOCX 文档")
            main = etree.fromstring(archive.read("word/document.xml"), parser)
            word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            if (
                main.tag != f"{{{word_namespace}}}document"
                or main.find(f"{{{word_namespace}}}body") is None
            ):
                reject("DOCX 缺少有效的 Word 正文结构")
            complex_content = main.xpath(
                "//*[local-name()='drawing' or local-name()='pict' "
                "or local-name()='txbxContent' or local-name()='altChunk']"
            )
            doc = Document(io.BytesIO(data))
    except HTTPException:
        raise
    except (
        zipfile.BadZipFile,
        KeyError,
        ValueError,
        OSError,
        etree.XMLSyntaxError,
        RuntimeError,
        NotImplementedError,
    ):
        reject("DOCX 文件损坏或格式不支持")
    issues = (
        [issue("图片、图形和复杂文本框未提取，请改用正文或普通表格", severity="warning")]
        if complex_content
        else []
    )
    if kind == "notes":
        rows = [
            {"location": f"段落 {i}", "values": {"text": p.text}}
            for i, p in enumerate(doc.paragraphs, 1)
            if p.text.strip()
        ]
        for ti, table in enumerate(doc.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                rows.append(
                    {
                        "location": f"表格 {ti} / 行 {ri}",
                        "values": {"text": " | ".join(c.text for c in row.cells)},
                    }
                )
        if not rows:
            reject("DOCX 没有可提取的正文或普通表格")
        if len(rows) > MAX_ROWS:
            reject("DOCX 正文和表格超过 20000 行，请拆分文件")
        return ["text"], rows, issues
    if not doc.tables:
        reject("结构化材料需要普通表格；正文说明请选择文字材料分类")
    columns, rows, table_issues = [], [], []
    for ti, table in enumerate(doc.tables, 1):
        found_columns, found_rows, found_issues = table_rows(
            [
                (f"表格 {ti} / 行 {ri}", [c.text for c in row.cells])
                for ri, row in enumerate(table.rows, 1)
            ]
        )
        if columns and found_columns != columns:
            reject("同一结构化 DOCX 的各表格必须使用相同表头", f"表格 {ti}")
        columns = found_columns
        rows.extend(found_rows)
        table_issues.extend(found_issues)
        if len(rows) > MAX_ROWS:
            reject("明细超过 20000 行，请拆分文件")
    return columns, rows, issues + table_issues


def parse_material(data, name, kind):
    suffix = Path(name).suffix.lower()
    if suffix == ".csv":
        return parse_csv(data)
    if suffix == ".docx":
        return parse_docx(data, kind)
    if suffix == ".txt":
        if kind != "notes":
            reject("TXT 是文字说明材料，请选择文字材料分类")
        rows = [
            {"location": f"行 {i}", "values": {"text": line}}
            for i, line in enumerate(decode_text(data).splitlines(), 1)
            if line.strip()
        ]
        if not rows:
            reject("文字材料为空")
        if len(rows) > MAX_ROWS:
            reject("文字超过 20000 行，请拆分文件")
        return ["text"], rows, []
    reject("仅支持 TXT、CSV、DOCX 文件")


def suggest_mapping(kind, columns):
    result = {}
    for field in FIELDS[kind] + OPTIONAL_FIELDS[kind]:
        candidates = [field, *ALIASES.get(field, [])]
        matches = [c for c in columns if c.casefold() in candidates]
        if len(matches) == 1:
            result[field] = matches[0]
    return result


def preview_rows(rows):
    return [
        {"location": row["location"], "values": {k: v[:1000] for k, v in row["values"].items()}}
        for row in rows[:20]
    ]


def validate_rows(kind, columns, rows, mapping):
    issues = []
    if kind == "notes":
        if mapping:
            issues.append(issue("文字材料不使用字段映射", "字段映射"))
        return issues
    for key in FIELDS[kind]:
        if key not in mapping or mapping[key] not in columns:
            issues.append(issue(f"请映射必需字段 {key}", "字段映射"))
    if any(value not in columns for value in mapping.values()) or any(
        key not in FIELDS[kind] + OPTIONAL_FIELDS[kind] for key in mapping
    ):
        issues.append(issue("映射含未知字段或不存在的原始列", "字段映射"))
    if len(set(mapping.values())) != len(mapping):
        issues.append(issue("一个原始列不能映射到多个字段", "字段映射"))
    if issues:
        return issues
    seen = set()
    for row in rows:
        values = {key: row["values"].get(column, "") for key, column in mapping.items()}
        location = row["location"]
        for key in FIELDS[kind]:
            value = values[key]
            if not value:
                severity = (
                    "warning"
                    if kind == "attendance" and key in {"planned_hours", "actual_hours"}
                    else "error"
                )
                issues.append(issue(f"{key} 未提供，不能当作 0", location, severity))
        for key, value in values.items():
            if key == "valid" and value.casefold() not in {"true", "false", "1", "0"}:
                issues.append(issue("valid 须为 true/false/1/0", location))
            if (key == "date" or key.endswith("_date")) and value:
                try:
                    if not valid_date(value):
                        raise ValueError
                except ValueError:
                    issues.append(issue(f"{key} 须为有效 YYYY-MM-DD 日期", location))
            if (
                key
                in {
                    "progress",
                    "planned_hours",
                    "actual_hours",
                    "score",
                    "scale_min",
                    "scale_max",
                    "leave_hours",
                    "absent_hours",
                    "late_minutes",
                }
                and value
            ):
                try:
                    number = float(value)
                    if (
                        not math.isfinite(number)
                        or (key not in {"score", "scale_min", "scale_max"} and number < 0)
                        or (key == "progress" and number > 100)
                        or (key.endswith("_hours") and number > 24)
                    ):
                        raise ValueError
                except ValueError:
                    issues.append(issue(f"{key} 数值无效或超出范围", location))
        if kind == "satisfaction":
            try:
                low, high, score = (float(values[k]) for k in ("scale_min", "scale_max", "score"))
                if low >= high or not low <= score <= high:
                    issues.append(issue("评分不在有效量表范围内", location))
            except ValueError:
                pass
        if (
            kind == "projects"
            and values["effective_date"]
            and values["snapshot_date"]
            and values["snapshot_date"] < values["effective_date"]
        ):
            issues.append(issue("快照日期不能早于生效日期", location))
        keys = {
            "projects": ("project_id", "snapshot_date"),
            "attendance": ("person_id", "date"),
            "achievements": ("achievement_id",),
            "satisfaction": ("feedback_id",),
        }[kind]
        identity = tuple(values[k] for k in keys)
        if identity in seen:
            issues.append(issue("重复业务记录：" + "+".join(keys), location))
        seen.add(identity)
        if len(issues) >= 200:
            issues.append(issue("问题过多，仅显示前 200 项；请修正后重新导入"))
            break
    return issues


class MaterialService:
    def __init__(self, connection, material_root: Path):
        self.connection = connection
        self.root = Path(material_root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        with connection() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS imports (
                    id TEXT PRIMARY KEY, workspace TEXT NOT NULL, payload TEXT NOT NULL,
                    relative_path TEXT NOT NULL, rows_json TEXT NOT NULL,
                    confirmed_material_id TEXT, confirmation_hash TEXT
                );
                CREATE TABLE IF NOT EXISTS material_versions (
                    id TEXT PRIMARY KEY, workspace TEXT NOT NULL, name TEXT NOT NULL,
                    kind TEXT NOT NULL, version INTEGER NOT NULL, fingerprint TEXT NOT NULL,
                    relative_path TEXT NOT NULL, payload TEXT NOT NULL, rows_json TEXT NOT NULL,
                    UNIQUE(workspace, fingerprint), UNIQUE(workspace,name,kind,version)
                );
            """)

    def stage_import(
        self,
        data: bytes,
        name: str,
        kind: str,
        workspace: str,
        source: str = "file",
        department="",
        start_date="",
        end_date="",
    ):
        if (
            workspace not in {"normal", "demo"}
            or kind not in FIELDS
            or source not in {"file", "local_api", "demo"}
        ):
            reject("材料分类、工作区或来源无效")
        if source == "demo" and workspace != "demo":
            reject("演示案例来源仅能导入演示空间", status=403)
        if (
            not name
            or len(name) > 200
            or name in {".", ".."}
            or any(c in name for c in "/\\\x00")
            or any(ord(c) < 32 for c in name)
        ):
            reject("文件名无效，请使用不含路径的普通文件名")
        if not data or len(data) > MAX_BYTES:
            reject("文件为空或超过 10MiB，请拆分后导入", status=413)
        if len(department) > 80 or len(start_date) > 10 or len(end_date) > 10:
            reject("部门或期间过长")
        columns, rows, issues = parse_material(data, name, kind)
        draft_id = uuid4().hex
        relative = f"{workspace}/{draft_id}{Path(name).suffix.lower()}"
        target = self.root / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        if not target.resolve().is_relative_to(self.root):
            reject("材料存储目录无效")
        payload = dict(
            id=draft_id,
            workspace=workspace,
            name=name,
            kind=kind,
            source=source,
            sha256=hashlib.sha256(data).hexdigest(),
            byte_size=len(data),
            department=department.strip(),
            start_date=start_date.strip(),
            end_date=end_date.strip(),
            columns=columns,
            suggested_mapping=suggest_mapping(kind, columns),
            preview=preview_rows(rows),
            row_count=len(rows),
            issues=issues,
            confirmed_material_id=None,
        )
        try:
            with target.open("xb") as handle:
                handle.write(data)
            target.chmod(0o600)
            with self.connection() as db:
                db.execute(
                    "INSERT INTO imports(id,workspace,payload,relative_path,rows_json) "
                    "VALUES(?,?,?,?,?)",
                    (draft_id, workspace, canonical_json(payload), relative, canonical_json(rows)),
                )
        except Exception:
            target.unlink(missing_ok=True)
            raise
        return payload

    @staticmethod
    def _prepare_confirmation(row, request):
        draft = json.loads(row["payload"])
        metadata = {
            key: (
                getattr(request, key) if getattr(request, key) is not None else draft[key]
            ).strip()
            for key in ("department", "start_date", "end_date")
        }
        canonical = dict(mapping=request.mapping, **metadata)
        digest = hashlib.sha256(canonical_json(canonical).encode()).hexdigest()
        issues = list(draft["issues"])
        for key in ("start_date", "end_date"):
            if metadata[key]:
                try:
                    if not valid_date(metadata[key]):
                        raise ValueError
                except ValueError:
                    issues.append(issue(f"{key} 须为有效 YYYY-MM-DD 日期", "范围"))
        if (
            metadata["start_date"]
            and metadata["end_date"]
            and metadata["end_date"] <= metadata["start_date"]
        ):
            issues.append(issue("结束日必须晚于开始日，统计不包含结束日", "范围"))
        if draft["kind"] == "notes" and not all(metadata.values()):
            issues.append(issue("文字材料须确认部门和起止期间", "范围"))
        rows = json.loads(row["rows_json"])
        issues.extend(validate_rows(draft["kind"], draft["columns"], rows, request.mapping))
        return draft, metadata, digest, issues

    def validate_import(self, import_id, workspace, payload):
        """Read-only preflight using exactly the same rules as confirmation."""
        request = (
            payload if isinstance(payload, ConfirmInput) else ConfirmInput.model_validate(payload)
        )
        with self.connection() as db:
            row = db.execute(
                "SELECT * FROM imports WHERE id=? AND workspace=?", (import_id, workspace)
            ).fetchone()
        if row is None:
            raise HTTPException(404, "未找到该工作区的导入草稿")
        _, _, digest, issues = self._prepare_confirmation(row, request)
        if row["confirmed_material_id"] and row["confirmation_hash"] != digest:
            raise HTTPException(409, "该导入已按其他配置确认，请重新导入以建立新版本")
        return issues[:201]

    def confirm(self, import_id, workspace, payload):
        request = (
            payload if isinstance(payload, ConfirmInput) else ConfirmInput.model_validate(payload)
        )
        with self.connection() as db:
            db.execute("BEGIN IMMEDIATE")
            row = db.execute(
                "SELECT * FROM imports WHERE id=? AND workspace=?", (import_id, workspace)
            ).fetchone()
            if row is None:
                raise HTTPException(404, "未找到该工作区的导入草稿")
            draft, metadata, digest, issues = self._prepare_confirmation(row, request)
            canonical = dict(mapping=request.mapping, **metadata)
            if row["confirmed_material_id"]:
                if row["confirmation_hash"] != digest:
                    raise HTTPException(409, "该导入已按其他配置确认，请重新导入以建立新版本")
                return json.loads(
                    db.execute(
                        "SELECT payload FROM material_versions WHERE id=?",
                        (row["confirmed_material_id"],),
                    ).fetchone()[0]
                )
            if any(item["severity"] == "error" for item in issues):
                raise HTTPException(
                    422, {"message": "材料校验未通过，尚未入库", "issues": issues[:201]}
                )
            fingerprint = hashlib.sha256(
                canonical_json(
                    dict(sha256=draft["sha256"], kind=draft["kind"], **canonical)
                ).encode()
            ).hexdigest()
            existing = db.execute(
                "SELECT * FROM material_versions WHERE workspace=? AND fingerprint=?",
                (workspace, fingerprint),
            ).fetchone()
            if existing:
                material = json.loads(existing["payload"])
            else:
                version = db.execute(
                    "SELECT COALESCE(MAX(version),0)+1 FROM material_versions "
                    "WHERE workspace=? AND name=? AND kind=?",
                    (workspace, draft["name"], draft["kind"]),
                ).fetchone()[0]
                material = {
                    key: value
                    for key, value in draft.items()
                    if key not in {"suggested_mapping", "confirmed_material_id"}
                }
                material.update(
                    id=uuid4().hex,
                    version=version,
                    mapping=request.mapping,
                    issues=issues,
                    created_at=datetime.now(UTC).isoformat(),
                    **metadata,
                )
                db.execute(
                    "INSERT INTO material_versions VALUES(?,?,?,?,?,?,?,?,?)",
                    (
                        material["id"],
                        workspace,
                        material["name"],
                        material["kind"],
                        version,
                        fingerprint,
                        row["relative_path"],
                        canonical_json(material),
                        row["rows_json"],
                    ),
                )
            db.execute(
                "UPDATE imports SET confirmed_material_id=?, confirmation_hash=? WHERE id=?",
                (material["id"], digest, import_id),
            )
            return material

    def get(self, material_id, workspace):
        with self.connection() as db:
            row = db.execute(
                "SELECT * FROM material_versions WHERE id=? AND workspace=?",
                (material_id, workspace),
            ).fetchone()
        if row is None:
            raise HTTPException(404, "未找到该工作区中的材料")
        return dict(row)

    def original_path(self, row):
        target = (self.root / row["relative_path"]).resolve()
        if not target.is_relative_to(self.root) or not target.is_file():
            raise HTTPException(404, "材料原件不可用")
        return target


def create_material_router(connection, material_root: Path, demo_root: Path):
    router = APIRouter()
    service = MaterialService(connection, material_root)
    router.material_service = service

    @router.post("/api/v1/materials/import", status_code=201)
    async def import_material(
        file: Annotated[UploadFile, File()],
        kind: Annotated[Kind, Form()],
        department: Annotated[str, Form(max_length=80)] = "",
        start_date: Annotated[str, Form(max_length=10)] = "",
        end_date: Annotated[str, Form(max_length=10)] = "",
        x_workspace: Annotated[Workspace, Header()] = "normal",
    ):
        try:
            data = await file.read(MAX_BYTES + 1)
            return service.stage_import(
                data,
                file.filename or "",
                kind,
                x_workspace,
                department=department,
                start_date=start_date,
                end_date=end_date,
            )
        finally:
            await file.close()

    @router.post("/api/v1/imports/{import_id}/confirm")
    def confirm_material(
        import_id: str,
        payload: ConfirmInput,
        x_workspace: Annotated[Workspace, Header()] = "normal",
    ):
        return service.confirm(import_id, x_workspace, payload)

    @router.get("/api/v1/materials")
    def list_materials(
        x_workspace: Annotated[Workspace, Header()] = "normal",
        search: Annotated[str, Query(max_length=200)] = "",
        kind: Kind | Literal[""] = "",
    ):
        with connection() as db:
            rows = db.execute(
                "SELECT payload FROM material_versions WHERE workspace=? "
                "AND instr(lower(name || char(10) || kind || char(10) || "
                "COALESCE(json_extract(payload,'$.department'),'')),lower(?))>0 "
                "AND (?='' OR kind=?) ORDER BY rowid DESC",
                (x_workspace, search.strip(), kind, kind),
            ).fetchall()
        return {"items": [json.loads(row[0]) for row in rows]}

    @router.get("/api/v1/materials/{material_id}")
    def material_detail(material_id: str, x_workspace: Annotated[Workspace, Header()] = "normal"):
        return json.loads(service.get(material_id, x_workspace)["payload"])

    @router.get("/api/v1/materials/{material_id}/download")
    def download_material(material_id: str, x_workspace: Annotated[Workspace, Header()] = "normal"):
        row = service.get(material_id, x_workspace)
        return FileResponse(
            service.original_path(row), filename=row["name"], media_type="application/octet-stream"
        )

    return router
